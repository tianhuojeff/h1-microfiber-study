from pathlib import Path
from copy import deepcopy
import json
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
TEMPLATE = ROOT / "template_originals" / "专利检索分析报告模板_参赛选手信用承诺书_官方原件.docx"
OUTPUT = ROOT / "H1专利检索分析报告_官方模板草稿_20260928.docx"
CONTENT = ROOT / "H1专利检索分析参赛报告_内容稿_20260928.md"
FIGURE = ROOT / "图1_冻结同族_流路与移出证据标签_20260928.png"
MATRIX = ROOT.parent / "patents" / "training_20260927" / "14同族人工功能矩阵.json"
ARCHIVED_RECORD = ROOT / "search_screenshots" / "google_patents_archived_public_record_WO2021116933A1_20260928.png"
ARCHIVED_ALSO_PUBLISHED = ROOT / "search_screenshots" / "archived_also_published_as_readable_20260928.png"

def set_font(run, name="仿宋_GB2312", size=12, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)

def set_spacing(paragraph, before=0, after=6, line=1.45):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line

def add_text(doc, text, size=12, bold=False, align=None, first_indent=True, before=0, after=6):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if first_indent:
        p.paragraph_format.first_line_indent = Pt(24)
    set_spacing(p, before, after)
    r = p.add_run(text)
    set_font(r, size=size, bold=bold)
    return p

def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    set_spacing(p, before=12 if level == 1 else 8, after=5, line=1.2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_font(r, name="黑体", size=15 if level == 1 else 13, bold=True)
    return p

def clear_body_keep_section(doc):
    body = doc._element.body
    sectPr = body.sectPr
    for child in list(body):
        if child is not sectPr:
            body.remove(child)

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), color)
    tcPr.append(shd)

def cell_text(cell, text, bold=False, size=10.5, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    set_spacing(p, after=0, line=1.15)
    r = p.add_run(text)
    set_font(r, size=size, bold=bold)
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement('w:tcMar')
    for edge in ('top', 'start', 'bottom', 'end'):
        node = OxmlElement(f'w:{edge}')
        node.set(qn('w:w'), '90')
        node.set(qn('w:type'), 'dxa')
        mar.append(node)
    tcPr.append(mar)

def add_search_table(doc):
    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    table.style = 'Table Grid'
    widths = [Inches(1.35), Inches(2.55), Inches(2.85)]
    for i, text in enumerate(["检索字段", "中文", "英文 / 分类说明"]):
        cell_text(table.rows[0].cells[i], text, bold=True, center=True)
        shade(table.rows[0].cells[i], 'D9E2F3')
        table.rows[0].cells[i].width = widths[i]
    rows = [
        ("关键词", "洗衣；洗衣机；微纤维；微塑料；过滤；收集；可拆；清理/再生", "laundry; washing machine; microfiber; microfibre; microplastic; filter; collection; removable; self-cleaning"),
        ("实际窄题名式", "TI=((laundry OR \"washing machine\") (microfiber OR microfibre OR microplastic))", "题名字段；未限定申请人、地域或申请/公开时间"),
        ("分类号", "本次冻结样本未以 IPC 分类号作为纳入门槛", "分类号仅作后续逐件核验辅助；不以模板示例的 H01L 31/04 作为本报告分类"),
    ]
    for left, mid, right in rows:
        cells = table.add_row().cells
        for i, value in enumerate((left, mid, right)):
            cell_text(cells[i], value, size=10.5, center=(i == 0))
            cells[i].width = widths[i]
    p = doc.add_paragraph()
    set_spacing(p, after=5)

def split_markdown(content):
    sections = {}
    current = None
    for raw in content.splitlines():
        line = raw.strip()
        match = re.match(r'^##\s+(.+)$', line)
        if match:
            current = match.group(1)
            sections[current] = []
        elif current is not None:
            sections[current].append(raw.rstrip())
    return sections

def is_placeholder(line):
    return '检索截图待补位置' in line

def add_markdown_lines(doc, lines, allow_figure=False):
    for raw in lines:
        line = raw.strip()
        if not line or line.startswith('> 使用说明'):
            continue
        if '图1插入位置' in line or '图1图注建议' in line:
            continue
        if line.startswith('### '):
            add_heading(doc, line[4:], level=2)
            continue
        line = re.sub(r'^\*\*(.+)\*\*$', r'\1', line)
        if is_placeholder(line):
            if '待补位置1' in line:
                add_text(doc, 'G1当时未保留现场页面截图，后续访问触发GoogleSorry限流，无法补拍。下列两图为已保存公开页快照的离线回放，仅辅助说明F1定向扩展线索，不作为G1题名检索式、98即时计数或现场数据库页面的证明。', size=10.5, first_indent=True, before=6, after=6)
                continue
            if '待补位置2' in line and ARCHIVED_RECORD.exists() and ARCHIVED_ALSO_PUBLISHED.exists():
                add_text(doc, '图2a—图2b为已保存Google Patents公开页快照的离线浏览器回放，仅辅助说明WO2021116933A1题录及其“Also Published As”栏所列的CN115087774A版本线索；不证明现场数据库查询、题名检索式、98即时计数、法律同族、权利范围或法律状态。', size=10.5, first_indent=True, before=6, after=4)
                for image_path, caption in (
                    (ARCHIVED_RECORD, '图2a 已保存公开页快照的离线回放：WO2021116933A1题录区'),
                    (ARCHIVED_ALSO_PUBLISHED, '图2b 已保存公开页快照的离线回放：“Also Published As”栏含CN115087774A版本线索'),
                ):
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    set_spacing(p, before=2, after=2)
                    p.add_run().add_picture(str(image_path), width=Inches(5.8))
                    add_text(doc, caption, size=9.5, first_indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, after=5)
                continue
            add_text(doc, '【待补真实截图：' + re.sub(r'^【|】$', '', line).replace('检索截图待补位置1：', '').replace('检索截图待补位置2：', '') + '】', size=11, bold=True, first_indent=False, before=6, after=6)
            continue
        line = line.replace('**', '')
        numbered = re.match(r'^\d+\.\s+(.+)$', line)
        if numbered:
            add_text(doc, numbered.group(1), size=12, first_indent=True)
            continue
        if line.startswith('|') or line.startswith('---'):
            continue
        add_text(doc, line, size=12)

def add_sources(doc, matrix_rows):
    add_heading(doc, '参考资料与公开件来源目录', level=1)
    add_text(doc, '以下链接对应冻结矩阵中14个代表公开文本的 source 字段，访问与核验日期为2026年9月27日至28日。公开文本用于技术信息比较，不构成法律状态、有效性、侵权或自由实施意见。', size=11, first_indent=True)
    add_text(doc, '本地论文全文共14篇，仅用于选题背景和研究边界，不计入专利同族分母，正文未据此引用性能数。已核的背景文献包括：', size=11, first_indent=True, before=4)
    add_text(doc, 'McIlwraith, H. K. et al. (2019). Capturing microfibers - marketed technologies reduce microfiber emissions from washing machines. Marine Pollution Bulletin. DOI: 10.1016/j.marpolbul.2018.12.012。', size=10, first_indent=True, after=3)
    add_text(doc, 'Erdle, L. M. et al. (2021). Washing Machine Filters Reduce Microfiber Emissions: Evidence From a Community-Scale Pilot in Parry Sound, Ontario. Frontiers in Marine Science. DOI: 10.3389/fmars.2021.777865。', size=10, first_indent=True, after=5)
    for row in matrix_rows:
        add_text(doc, f"{row['publication']}：{row['source']}", size=9.5, first_indent=False, after=2)
    add_text(doc, '本地统计底表《冻结样本数据说明与统计表_20260928》、检索日志《检索过程与样本形成记录_20260928》及《结论与亮点证据索引_20260928》为内部复核资料，用于归并依据、统计口径、检索式和权项层级回查；不作为公开专利来源的替代。', size=11, first_indent=True, before=5)
    add_text(doc, '数据限制：所有“x/14”和百分比均为冻结定向样本的多标签文本披露频次；不表示领域总体、市场份额、地域格局、法律状态、产品性能、环境减排效果或最终处置事实。', size=11, first_indent=True)

def main():
    content = CONTENT.read_text(encoding='utf-8')
    # The source manuscript remains immutable; this delivery corrects the current local-paper count.
    content = content.replace('11篇论文、待核线索', '14篇本地论文全文、待核线索')
    sections = split_markdown(content)
    rows = json.loads(MATRIX.read_text(encoding='utf-8'))
    doc = Document(TEMPLATE)
    clear_body_keep_section(doc)
    for section in doc.sections:
        section.top_margin = Inches(0.85)
        section.bottom_margin = Inches(0.78)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    p = doc.add_paragraph()
    set_spacing(p, after=8)
    r = p.add_run('附件1')
    set_font(r, name='黑体', size=14, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(p, after=4, line=1.2)
    r = p.add_run('2026年在京高校知识产权信息检索大赛检索实务赛')
    set_font(r, name='黑体', size=16, bold=False)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(p, after=14, line=1.2)
    r = p.add_run('专利检索分析报告')
    set_font(r, name='黑体', size=16, bold=False)
    add_text(doc, '草稿状态：待最终审核；本文件暂不作为最终提交版。', size=10.5, bold=True, first_indent=False, before=0, after=10)

    title = sections['一、报告名称'][1].strip().replace('**', '')
    add_heading(doc, '一、报告名称', level=1)
    add_text(doc, title, size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False, after=10)

    add_heading(doc, '二、选题背景', level=1)
    add_heading(doc, '1. 所选专利或领域', level=2)
    add_text(doc, '洗衣微纤维/微塑料收集装置及其拦截后的清理、转移和最终移出路径。', size=12, first_indent=True)
    add_heading(doc, '2. 目的与意义', level=2)
    add_markdown_lines(doc, sections['二、选题背景（所选领域、目的意义）'])

    add_heading(doc, '三、检索过程', level=1)
    add_heading(doc, '1. 数据来源', level=2)
    add_text(doc, 'Google Patents公开页面及其同族/国别页面；已保存公开页HTML、部分PDF和权利要求摘录；辅助网页搜索仅用于发现种子线索。', size=12)
    add_heading(doc, '2. 检索的时间范围', level=2)
    add_text(doc, 'G1检索式未设置申请日、优先权日或公开日过滤；实际查询与冻结核对日期为2026年9月28日（Asia/Shanghai）。冻结样本关联的优先权年份为2019—2024年，仅反映14个目的性样本的构成，不能解释为全领域的时间趋势。', size=12)
    add_heading(doc, '3. 检索的地域范围', level=2)
    add_text(doc, 'G1检索式未设置国家、地区或法域过滤。冻结矩阵代表公开文本中CN为10件、WO为4件，仅是本次代表文本前缀构成，不能作为地域份额、申请比例或竞争格局结论。', size=12)
    add_heading(doc, '4. 检索要素', level=2)
    add_search_table(doc)
    add_heading(doc, '5. 检索思路', level=2)
    add_markdown_lines(doc, sections['三、检索过程'])

    add_heading(doc, '四、分析过程及内容', level=1)
    add_markdown_lines(doc, sections['四、分析过程及内容'])
    if FIGURE.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_spacing(p, before=5, after=3)
        p.add_run().add_picture(str(FIGURE), width=Inches(5.9))
        add_text(doc, '图1 冻结同族样本的流路与移出证据标签（分母为14个同族；多标签计数）', size=10.5, first_indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)

    add_heading(doc, '五、结论建议', level=1)
    add_markdown_lines(doc, sections['五、结论建议'])
    add_heading(doc, '六、应用价值与亮点评述', level=1)
    add_markdown_lines(doc, sections['六、应用价值与亮点评述'])
    add_sources(doc, rows)

    # Remove any accidental personal metadata inherited from the source file.
    props = doc.core_properties
    props.author = ''
    props.last_modified_by = ''
    props.title = '洗衣微纤维专利检索分析报告'
    props.subject = '2026年在京高校知识产权信息检索大赛检索实务赛附件1草稿'
    props.comments = ''
    doc.save(OUTPUT)
    print(OUTPUT)

if __name__ == '__main__':
    main()
